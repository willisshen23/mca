# -*- coding: utf-8 -*-
"""生成《营销活动分析助手 · 核心技术介绍》三页 PPT —— 华为风格（丰富版）。"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ── 华为风格设计系统 ──────────────────────────────────────
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

RED      = RGBColor(0xC7, 0x00, 0x0B)   # 华为红
RED_DARK = RGBColor(0x9E, 0x00, 0x08)   # 深红
INK      = RGBColor(0x1A, 0x1A, 0x1A)   # 主文字（近黑）
GRAY     = RGBColor(0x59, 0x59, 0x59)   # 次级文字
LGRAY    = RGBColor(0x8C, 0x8C, 0x8C)   # 弱化文字
LINE     = RGBColor(0xE0, 0xE0, 0xE0)   # 分隔细线
BOX_BG   = RGBColor(0xF7, 0xF7, 0xF7)   # 浅灰底
WHITE    = RGBColor(0xFF, 0xFF, 0xFF)

FONT = "微软雅黑"


def set_run(run, text, size=14, bold=False, color=INK):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = FONT
    run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {})
        rPr.append(ea)
    ea.set("typeface", FONT)


def _fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def _rect(slide, x, y, w, h, color, shape=MSO_SHAPE.RECTANGLE, radius=None):
    sp = slide.shapes.add_shape(shape, x, y, w, h)
    _fill(sp, color)
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sp.adjustments[0] = radius
        except Exception:
            pass
    return sp


def _tb(slide, x, y, w, h):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    return tb, tf


def _shape_text(shape, text, size=12.5, bold=False, color=INK, align=PP_ALIGN.LEFT):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.12)
    tf.margin_right = Inches(0.1)
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = align
    set_run(p.add_run(), text, size=size, bold=bold, color=color)


def header(slide, title, subtitle=""):
    """华为风标题：红色大标题 + 红色短下划线。"""
    tb, tf = _tb(slide, Inches(0.55), Inches(0.34), Inches(10.5), Inches(0.62))
    set_run(tf.paragraphs[0].add_run(), title, size=26, bold=True, color=RED)
    _rect(slide, Inches(0.57), Inches(1.02), Inches(0.95), Inches(0.045), RED)
    if subtitle:
        tb2, tf2 = _tb(slide, Inches(0.57), Inches(1.14), Inches(11.5), Inches(0.32))
        set_run(tf2.paragraphs[0].add_run(), subtitle, size=12, color=LGRAY)


def footer(slide, page_no):
    _rect(slide, Inches(0.55), Inches(7.06), Inches(12.23), Inches(0.012), LINE)
    tb, tf = _tb(slide, Inches(0.55), Inches(7.15), Inches(6.0), Inches(0.3))
    p = tf.paragraphs[0]
    set_run(p.add_run(), "■ ", size=9, color=RED)
    set_run(p.add_run(), "营销活动分析助手", size=9, color=LGRAY)
    tb2, tf2 = _tb(slide, Inches(11.3), Inches(7.15), Inches(1.48), Inches(0.3))
    tf2.paragraphs[0].alignment = PP_ALIGN.RIGHT
    set_run(tf2.paragraphs[0].add_run(), page_no, size=9, color=LGRAY)


def value_bar(slide, text, y=Inches(6.28)):
    _rect(slide, Inches(0.55), y, Inches(12.23), Inches(0.62), BOX_BG)
    _rect(slide, Inches(0.55), y, Inches(0.07), Inches(0.62), RED)
    tb, tf = _tb(slide, Inches(0.85), y, Inches(11.7), Inches(0.62))
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    set_run(p.add_run(), "业务价值　", size=13, bold=True, color=RED)
    set_run(p.add_run(), text, size=13, color=INK)


def _card(slide, x, y, w, h, idx, title, lines,
          size=12.5, space_after=8, line_spacing=1.18):
    sp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    sp.fill.solid()
    sp.fill.fore_color.rgb = WHITE
    sp.adjustments[0] = 0.035
    sp.line.color.rgb = LINE
    sp.line.width = Pt(1.0)
    dot = _rect(slide, x + Inches(0.3), y + Inches(0.3), Inches(0.52), Inches(0.52), RED, MSO_SHAPE.OVAL)
    dtf = dot.text_frame
    dtf.word_wrap = False
    dtf.margin_left = 0
    dtf.margin_right = 0
    dtf.margin_top = 0
    dtf.margin_bottom = 0
    dtf.vertical_anchor = MSO_ANCHOR.MIDDLE
    dp = dtf.paragraphs[0]
    dp.alignment = PP_ALIGN.CENTER
    set_run(dp.add_run(), idx, size=15, bold=True, color=WHITE)
    tb, tf = _tb(slide, x + Inches(1.0), y + Inches(0.28), w - Inches(1.3), Inches(0.6))
    set_run(tf.paragraphs[0].add_run(), title, size=17, bold=True, color=INK)
    _rect(slide, x + Inches(0.3), y + Inches(1.0), w - Inches(0.6), Inches(0.014), LINE)
    tb2, tf2 = _tb(slide, x + Inches(0.3), y + Inches(1.18), w - Inches(0.6), h - Inches(1.4))
    first = True
    for line in lines:
        p = tf2.paragraphs[0] if first else tf2.add_paragraph()
        first = False
        p.space_after = Pt(space_after)
        p.line_spacing = line_spacing
        set_run(p.add_run(), "▪ ", size=size - 1, bold=True, color=RED)
        set_run(p.add_run(), line, size=size, color=INK)


def _arch_layer(slide, y, label, desc):
    """架构图单层：红色标签块 + 白底细边框内容区。"""
    h = Inches(0.62)
    lab = _rect(slide, Inches(0.55), y, Inches(1.9), h, RED)
    _shape_text(lab, label, size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(2.45), y, Inches(10.33), h)
    box.fill.solid()
    box.fill.fore_color.rgb = WHITE
    box.line.color.rgb = LINE
    box.line.width = Pt(1.0)
    _shape_text(box, desc, size=12.5, color=INK, align=PP_ALIGN.LEFT)


def _arch_arrow(slide, y, label):
    """层间向下箭头 + 说明文字。"""
    arr = _rect(slide, Inches(6.55), y, Inches(0.24), Inches(0.16), LGRAY, MSO_SHAPE.DOWN_ARROW)
    tb, tf = _tb(slide, Inches(6.95), y - Inches(0.03), Inches(3.0), Inches(0.22))
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    set_run(tf.paragraphs[0].add_run(), label, size=10, color=LGRAY)


prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H
blank = prs.slide_layouts[6]

# ══════════════════════════════════════════════════════════
# 第 1 页：项目定位与系统架构
# ══════════════════════════════════════════════════════════
s = prs.slides.add_slide(blank)
header(s, "营销活动分析助手", "面向银行营销场景的对话式智能分析平台")
footer(s, "01 / 03")

pos = _rect(s, Inches(0.55), Inches(1.62), Inches(12.23), Inches(1.0), BOX_BG)
_rect(s, Inches(0.55), Inches(1.62), Inches(0.07), Inches(1.0), RED)
tb, tf = _tb(s, Inches(0.88), Inches(1.76), Inches(11.7), Inches(0.76))
p = tf.paragraphs[0]
set_run(p.add_run(), "一句话定位　", size=12.5, bold=True, color=RED)
p2 = tf.add_paragraph()
p2.line_spacing = 1.2
set_run(p2.add_run(),
        "用自然语言对话，AI 自主完成营销活动检索、数据拉取、分析报告生成与导出；关键写操作由人工审批把关，全程可审计。",
        size=14, color=INK)

# 分层架构图（四层 + 向下数据流）
_arch_layer(s, Inches(2.82), "用户交互层",
            "Vue 3 前端 — 对话界面 · 审批卡片 · 报告预览 · 技能/模板/历史/日志管理")
_arch_arrow(s, Inches(3.46), "HTTP / SSE 流式")
_arch_layer(s, Inches(3.64), "应用服务层",
            "FastAPI 后端 — REST 接口 · SSE 流式推送 · 会话/技能/日志等路由")
_arch_arrow(s, Inches(4.28), "run / resume 调度")
_arch_layer(s, Inches(4.46), "智能体层",
            "agent_v3（openjiuwen）— 多模型路由 · 12+ 业务工具 · 报告生成管线 · 会话持久化")
_arch_arrow(s, Inches(5.10), "HTTP 接口调用")
_arch_layer(s, Inches(5.28), "数据服务层",
            "银行营销外部 API — 活动搜索 · 模板清单 · 组件数值 · 外呼活动数据（开发态由 mock 模拟）")

value_bar(s, "分析效率显著提升 · 报告专业规范 · 关键操作合规可控 · 可离线落地银行内网")

# ══════════════════════════════════════════════════════════
# 第 2 页：技术亮点一 —— 自主智能体与人机协同
# ══════════════════════════════════════════════════════════
s = prs.slides.add_slide(blank)
header(s, "技术亮点一 · 自主智能体与人机协同审批")
footer(s, "02 / 03")

_card(s, Inches(0.55), Inches(1.6), Inches(6.05), Inches(4.55), "1", "自主智能体（Agent）", [
    "ReAct 循环：大模型自主完成“思考 → 行动 → 观察”，多步任务自动拆解、动态调整执行路径",
    "12+ 业务工具：活动检索、模板/组件查询、数据解析、报告生成·修改·导出，端到端闭环",
    "多模型分层路由：主推理模型 + 轻量模型 + 降级备选，兼顾质量、成本与高可用",
    "任务规划辅助：内置待办工具，自动拆分并跟踪分析步骤",
    "安全护栏：每轮工具调用限次、长耗时任务超时豁免，防止失控与空转",
    "流式交互：任务进度实时可见，支持取消与断线重连回放",
], size=12, space_after=6, line_spacing=1.15)

_card(s, Inches(6.85), Inches(1.6), Inches(5.93), Inches(4.55), "2", "人机协同审批（HITL）", [
    "硬中断机制：写操作触发中断，智能体暂停等待人工决策，前端实时弹出审批卡片",
    "三种决策：批准执行 / 修改参数后执行 / 拒绝并反馈模型",
    "审批持久化：刷新或断线后可恢复，决策现场不丢失",
    "二次审批：拒绝后模型重试会再次触发审批，每一步都在掌控中",
    "全程审计留痕：会话日志、运行历史、报告版本链，可逐轮追溯",
    "敏感信息隔离：用户/机构/渠道等参数不入模型，仅经受控上下文注入",
], size=12, space_after=6, line_spacing=1.15)

value_bar(s, "让 AI 高效干活，让关键动作合规可控，人机各司其职", y=Inches(6.35))

# ══════════════════════════════════════════════════════════
# 第 3 页：技术亮点二 —— 智能报告管线与工程保障
# ══════════════════════════════════════════════════════════
s = prs.slides.add_slide(blank)
header(s, "技术亮点二 · 智能报告生成管线与工程保障")
footer(s, "03 / 03")

_card(s, Inches(0.55), Inches(1.6), Inches(6.05), Inches(4.55), "1", "智能报告生成管线", [
    "模板智能匹配：按活动类型与视角（外呼 / 短信 / 综合运营）自动匹配分析模板与组件",
    "组件数值拉取：对接银行营销接口自动取数，支持上传数据解析、外呼数据与支行下钻",
    "数据清洗与合成：数值去重归一后，由大模型生成结构化 HTML 报告与图表",
    "质量守卫：自动校验乱码 / 格式残留 / 缺失等问题，不合格自动重试修正",
    "版式守卫：检测 ECharts 布局反模式，确保图表渲染正常",
    "灵活可配：指标 / 维度 / 视角自由组合，报告多版本可回溯",
], size=12, space_after=6, line_spacing=1.15)

_card(s, Inches(6.85), Inches(1.6), Inches(5.93), Inches(4.55), "2", "工程质量与落地保障", [
    "分层架构：api / services / domain / infrastructure，职责清晰、易维护、可测试",
    "会话持久化：多轮上下文记忆，支持 SQLite / MySQL，断线重连不丢状态",
    "可观测审计：运行历史与报告版本链，问题可定位、过程可追溯",
    "工程防线：自动化测试覆盖智能体工具、审批与管线关键路径",
    "离线部署：支持内网与麒麟系统离线安装，启动期自动迁移建库",
    "模型多源兼容：兼容 OpenAI 协议，可接 GPUStack / vLLM / Ollama / DeepSeek 等",
], size=12, space_after=6, line_spacing=1.15)

value_bar(s, "报告专业可控、系统稳定可靠，满足银行内网合规与落地要求", y=Inches(6.35))

out = "营销活动分析助手_核心技术介绍.pptx"
prs.save(out)
print("已生成:", out)
