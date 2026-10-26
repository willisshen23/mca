# -*- coding: utf-8 -*-
"""生成《openJiuwen 框架应用实践》培训 PPT —— 华为风格，复用 generate_ppt.py 设计系统。"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ── 华为风格设计系统（与 generate_ppt.py 一致）──────────────
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

RED      = RGBColor(0xC7, 0x00, 0x0B)
RED_DARK = RGBColor(0x9E, 0x00, 0x08)
INK      = RGBColor(0x1A, 0x1A, 0x1A)
GRAY     = RGBColor(0x59, 0x59, 0x59)
LGRAY    = RGBColor(0x8C, 0x8C, 0x8C)
LINE     = RGBColor(0xE0, 0xE0, 0xE0)
BOX_BG   = RGBColor(0xF7, 0xF7, 0xF7)
WHITE    = RGBColor(0xFF, 0xFF, 0xFF)

FONT = "微软雅黑"
FOOTER_TEXT = "openJiuwen 框架应用实践"
TOTAL = 13


def set_run(run, text, size=14, bold=False, color=INK):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = FONT
    run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {})
        rPr.append(ea)
    ea.set("typeface", FONT)


def _fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def _rect(slide, x, y, w, h, color, shape=MSO_SHAPE.RECTANGLE, radius=None):
    sp = slide.shapes.add_shape(shape, x, y, w, h)
    _fill(sp, color)
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sp.adjustments[0] = radius
        except Exception:
            pass
    return sp


def _tb(slide, x, y, w, h):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    return tb, tf


def _shape_text(shape, text, size=12.5, bold=False, color=INK, align=PP_ALIGN.LEFT):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.12)
    tf.margin_right = Inches(0.1)
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = align
    set_run(p.add_run(), text, size=size, bold=bold, color=color)


def header(slide, title, subtitle=""):
    tb, tf = _tb(slide, Inches(0.55), Inches(0.34), Inches(10.5), Inches(0.62))
    set_run(tf.paragraphs[0].add_run(), title, size=26, bold=True, color=RED)
    _rect(slide, Inches(0.57), Inches(1.02), Inches(0.95), Inches(0.045), RED)
    if subtitle:
        tb2, tf2 = _tb(slide, Inches(0.57), Inches(1.14), Inches(11.5), Inches(0.32))
        set_run(tf2.paragraphs[0].add_run(), subtitle, size=12, color=LGRAY)


def footer(slide, page_no):
    _rect(slide, Inches(0.55), Inches(7.06), Inches(12.23), Inches(0.012), LINE)
    tb, tf = _tb(slide, Inches(0.55), Inches(7.15), Inches(6.0), Inches(0.3))
    p = tf.paragraphs[0]
    set_run(p.add_run(), "■ ", size=9, color=RED)
    set_run(p.add_run(), FOOTER_TEXT, size=9, color=LGRAY)
    tb2, tf2 = _tb(slide, Inches(11.3), Inches(7.15), Inches(1.48), Inches(0.3))
    tf2.paragraphs[0].alignment = PP_ALIGN.RIGHT
    set_run(tf2.paragraphs[0].add_run(), page_no, size=9, color=LGRAY)


def value_bar(slide, text, y=Inches(6.28), label="一句话总结"):
    _rect(slide, Inches(0.55), y, Inches(12.23), Inches(0.62), BOX_BG)
    _rect(slide, Inches(0.55), y, Inches(0.07), Inches(0.62), RED)
    tb, tf = _tb(slide, Inches(0.85), y, Inches(11.7), Inches(0.62))
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    set_run(p.add_run(), label + "　", size=13, bold=True, color=RED)
    set_run(p.add_run(), text, size=13, color=INK)


def _card(slide, x, y, w, h, idx, title, lines,
          size=12.5, space_after=8, line_spacing=1.18):
    sp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    sp.fill.solid()
    sp.fill.fore_color.rgb = WHITE
    sp.adjustments[0] = 0.035
    sp.line.color.rgb = LINE
    sp.line.width = Pt(1.0)
    dot = _rect(slide, x + Inches(0.3), y + Inches(0.3), Inches(0.52), Inches(0.52), RED, MSO_SHAPE.OVAL)
    dtf = dot.text_frame
    dtf.word_wrap = False
    dtf.margin_left = 0
    dtf.margin_right = 0
    dtf.margin_top = 0
    dtf.margin_bottom = 0
    dtf.vertical_anchor = MSO_ANCHOR.MIDDLE
    dp = dtf.paragraphs[0]
    dp.alignment = PP_ALIGN.CENTER
    set_run(dp.add_run(), idx, size=15, bold=True, color=WHITE)
    tb, tf = _tb(slide, x + Inches(1.0), y + Inches(0.28), w - Inches(1.3), Inches(0.6))
    set_run(tf.paragraphs[0].add_run(), title, size=17, bold=True, color=INK)
    _rect(slide, x + Inches(0.3), y + Inches(1.0), w - Inches(0.6), Inches(0.014), LINE)
    tb2, tf2 = _tb(slide, x + Inches(0.3), y + Inches(1.18), w - Inches(0.6), h - Inches(1.4))
    first = True
    for line in lines:
        p = tf2.paragraphs[0] if first else tf2.add_paragraph()
        first = False
        p.space_after = Pt(space_after)
        p.line_spacing = line_spacing
        set_run(p.add_run(), "▪ ", size=size - 1, bold=True, color=RED)
        set_run(p.add_run(), line, size=size, color=INK)


def _arch_layer(slide, y, label, desc, h=Inches(0.62)):
    lab = _rect(slide, Inches(0.55), y, Inches(1.9), h, RED)
    _shape_text(lab, label, size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(2.45), y, Inches(10.33), h)
    box.fill.solid()
    box.fill.fore_color.rgb = WHITE
    box.line.color.rgb = LINE
    box.line.width = Pt(1.0)
    _shape_text(box, desc, size=12.5, color=INK, align=PP_ALIGN.LEFT)


def _arch_arrow(slide, y, label):
    _rect(slide, Inches(6.55), y, Inches(0.24), Inches(0.16), LGRAY, MSO_SHAPE.DOWN_ARROW)
    tb, tf = _tb(slide, Inches(6.95), y - Inches(0.03), Inches(3.0), Inches(0.22))
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    set_run(tf.paragraphs[0].add_run(), label, size=10, color=LGRAY)


def _bullets(slide, x, y, w, h, lines, size=13, space_after=7, line_spacing=1.18, marker_color=RED):
    tb, tf = _tb(slide, x, y, w, h)
    first = True
    for line in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(space_after)
        p.line_spacing = line_spacing
        if isinstance(line, tuple):  # (加粗前缀, 正文)
            set_run(p.add_run(), "▪ ", size=size - 1, bold=True, color=marker_color)
            set_run(p.add_run(), line[0], size=size, bold=True, color=INK)
            set_run(p.add_run(), line[1], size=size, color=INK)
        else:
            set_run(p.add_run(), "▪ ", size=size - 1, bold=True, color=marker_color)
            set_run(p.add_run(), line, size=size, color=INK)


def _code_box(slide, x, y, w, h, code, size=11):
    box = _rect(slide, x, y, w, h, RGBColor(0x2B, 0x2B, 0x2B))
    tb, tf = _tb(slide, x + Inches(0.2), y + Inches(0.15), w - Inches(0.4), h - Inches(0.3))
    first = True
    for ln in code.split("\n"):
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.line_spacing = 1.12
        r = p.add_run()
        r.text = ln
        r.font.size = Pt(size)
        r.font.name = "Menlo"
        r.font.color.rgb = RGBColor(0xE8, 0xE8, 0xE8)


def _table(slide, x, y, w, h, rows, col_widths=None, header_row=True, size=11.5):
    tbl_shape = slide.shapes.add_table(len(rows), len(rows[0]), x, y, w, h)
    tbl = tbl_shape.table
    if col_widths:
        for i, cw in enumerate(col_widths):
            tbl.columns[i].width = cw
    for ri, row in enumerate(rows):
        for ci, cell_text in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.margin_left = Inches(0.08)
            cell.margin_top = Inches(0.03)
            cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tfp = cell.text_frame.paragraphs[0]
            is_head = header_row and ri == 0
            r = tfp.add_run()
            r.text = cell_text
            r.font.size = Pt(size)
            r.font.bold = is_head
            r.font.name = FONT
            r.font.color.rgb = WHITE if is_head else INK
            cell.fill.solid()
            cell.fill.fore_color.rgb = RED if is_head else (BOX_BG if ri % 2 == 0 else WHITE)


prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H
blank = prs.slide_layouts[6]


def page(no):
    return "%02d / %02d" % (no, TOTAL)


# ══════════════════════════════════════════════════════════
# 1. 封面
# ══════════════════════════════════════════════════════════
s = prs.slides.add_slide(blank)
_rect(s, 0, 0, SLIDE_W, SLIDE_H, WHITE)
_rect(s, 0, Inches(2.2), Inches(0.18), Inches(2.2), RED)
tb, tf = _tb(s, Inches(0.9), Inches(2.3), Inches(11.5), Inches(1.0))
set_run(tf.paragraphs[0].add_run(), "openJiuwen 框架应用实践", size=44, bold=True, color=INK)
tb, tf = _tb(s, Inches(0.9), Inches(3.35), Inches(11.5), Inches(0.6))
set_run(tf.paragraphs[0].add_run(), "从 langgraph / deepagents 到 openJiuwen agent-core 的迁移经验与踩坑总结",
        size=20, color=GRAY)
_rect(s, Inches(0.92), Inches(4.25), Inches(1.6), Inches(0.05), RED)
tb, tf = _tb(s, Inches(0.9), Inches(4.55), Inches(11.5), Inches(1.2))
p = tf.paragraphs[0]
p.space_after = Pt(6)
set_run(p.add_run(), "基于项目：营销活动分析助手（agent_v3）", size=14, color=INK)
p2 = tf.add_paragraph()
set_run(p2.add_run(), "框架版本 openjiuwen==0.1.16.post2 · Python 3.12 · 2026-08 完成迁移", size=12, color=LGRAY)
tb, tf = _tb(s, Inches(0.9), Inches(6.9), Inches(11.5), Inches(0.3))
set_run(tf.paragraphs[0].add_run(), "■ " + FOOTER_TEXT, size=9, color=LGRAY)

# ══════════════════════════════════════════════════════════
# 2. openJiuwen 是什么
# ══════════════════════════════════════════════════════════
s = prs.slides.add_slide(blank)
header(s, "openJiuwen 框架速览", "华为云 AgentArts 的开源版（2025-12 开源），Apache-2.0，Python ≥3.11")
footer(s, page(2))
_card(s, Inches(0.55), Inches(1.6), Inches(6.05), Inches(4.45), "1", "核心能力", [
    "ReActAgent：「思考 → 行动 → 观察」自主循环，Runner 驱动运行",
    "工具三形态：@tool 装饰器 / RestfulApi / MCPTool",
    "HITL：Rail 拦截器机制（不改 Agent 主流程实现人工审批）",
    "Workflow：组件化编排 + 条件分支，组件并发 / 流式传值内置",
    "多智能体：AgentAsTool（子 Agent 作工具）、TeamRuntime",
    "持久化：CheckpointerFactory 可插拔，内置 in_memory 与 PersistenceCheckpointer",
    "内置全链路 tracer（调测 / 可观测）",
], size=12, space_after=6, line_spacing=1.15)
_card(s, Inches(6.85), Inches(1.6), Inches(5.93), Inches(4.45), "2", "采用前必须知道的三件事", [
    "仍是 0.1.x：迭代极快（0.1.x 已发 25+ 版本）—— 钉版是硬性要求",
    "依赖体量大：全量 80 项（transformers / pymilvus / trafilatura 等），离线包要重做",
    "模型接入简单：OpenAI 兼容协议，api_base 可指任意端点（行内 deepseek 实测可用）",
    "文档覆盖不全：多处行为需实机 spike 验证（超时、迭代上限、异常包装等）",
], size=12, space_after=6, line_spacing=1.15)
value_bar(s, "政企同源、Apache-2.0、华为背书 —— 可持续性好，但 0.x 框架要用工程方法管住不确定性",
          y=Inches(6.35), label="选型结论")

# ══════════════════════════════════════════════════════════
# 3. 应用架构总览
# ══════════════════════════════════════════════════════════
s = prs.slides.add_slide(blank)
header(s, "本项目中的 agent_v3 架构", "openJiuwen 是唯一智能体运行时（v2 已删除，URL 前缀 /v2 仅为前端契约保留）")
footer(s, page(3))
_arch_layer(s, Inches(1.7), "前端", "Vue 3 — 对话界面 · 审批卡片 · 报告预览（SSE 帧契约不变，前端零改动）")
_arch_arrow(s, Inches(2.34), "HTTP / SSE")
_arch_layer(s, Inches(2.52), "API 层", "FastAPI — /run · /resume · /cancel · /runs/{id}/stream（断线重连游标重放）")
_arch_arrow(s, Inches(3.16), "run_agent_streaming")
_arch_layer(s, Inches(3.34), "运行时", "runtime.py — _BufferedRun 断线缓冲 · 取消泵 · resume 调度 · 审批登记")
_arch_arrow(s, Inches(3.98), "组装 / 拦截")
_arch_layer(s, Inches(4.16), "智能体", "ReActAgent 单例 + 5 个 Rail（审批 / 限额 / 幻觉拦截 / 进度帧）+ 13 个 @tool 业务工具")
_arch_arrow(s, Inches(4.80), "状态快照 / 恢复")
_arch_layer(s, Inches(4.98), "持久化", "PersistenceCheckpointer（SQLite WAL）续聊状态 + 业务表驱动历史重载（分层分离）")
value_bar(s, "全仓库只允许 agent_v3 包内 import openjiuwen —— 框架依赖收拢一处，业务代码零框架耦合",
          y=Inches(6.0), label="分层原则")

# ══════════════════════════════════════════════════════════
# 4. Agent 组装与工具注册
# ══════════════════════════════════════════════════════════
s = prs.slides.add_slide(blank)
header(s, "应用要点 ① · Agent 组装与工具注册", "backend/src/agent_v3/agent_factory.py")
footer(s, page(4))
_code_box(s, Inches(0.55), Inches(1.6), Inches(7.0), Inches(4.5), """client_cfg = ModelClientConfig(
    client_provider="OpenAI",   # OpenAI 兼容协议
    api_base=..., api_key=...,
    timeout=120, max_retries=..., verify_ssl=...)
request_cfg = ModelRequestConfig(model=..., temperature=...)

agent = ReActAgent(card=AgentCard(id=..., description=...))
    .configure(ReActAgentConfig(
        model_config_obj=request_cfg,
        model_client_config=client_cfg,
        prompt_template=[{"role": "system",
                          "content": build_system_prompt()}],
        max_iterations=50))     # 默认 5 太小，必须显式调大

for t in ALL_V3_TOOLS:          # 工具双注册
    Runner.resource_mgr.add_tool(t)
    agent.ability_manager.add(t.card)""")
_bullets(s, Inches(7.85), Inches(1.7), Inches(4.95), Inches(4.4), [
    ("进程级单例：", "asyncio.Lock 双重检查；工具只注册一次（重复注册打 ERROR 日志）"),
    ("max_iterations 默认 5：", "多工具编排下工具执行完就没迭代生成答复，直接报 error —— 显式调大到 50"),
    ("系统提示词：", "管理后台 DB 覆盖 + 8 条关键规则幂等追加"),
    ("LoggingModel：", "鸭子类型包装而非继承 Model —— 框架会探测能力方法，继承不调 super().__init__ 会 AttributeError"),
    ("Rail 注册顺序敏感：", "审批 → 幻觉拦截 → 限额 → 进度"),
], size=12, space_after=8)
value_bar(s, "组装期三件事：模型配置、工具双注册、Rail 按序注册；坑都标在注释里", y=Inches(6.35))

# ══════════════════════════════════════════════════════════
# 5. HITL 审批
# ══════════════════════════════════════════════════════════
s = prs.slides.add_slide(blank)
header(s, "应用要点 ② · HITL 审批：自定义 Rail 三决策", "backend/src/agent_v3/rails.py — ApprovalEditRail")
footer(s, page(5))
_code_box(s, Inches(0.55), Inches(1.6), Inches(7.0), Inches(4.5), """class ApprovalEditRail(BaseInterruptRail):
    async def resolve_interrupt(self, ctx, user_input):
        if user_input is None:          # 首次触发 → 挂起
            return self.interrupt(InterruptRequest(
                message=..., payload_schema=SCHEMA))
        decision = user_input.get("decision")
        if decision == "approve":
            return self.approve()
        if decision == "edit":          # 改写参数后放行
            return self.approve(new_args=json.dumps(edited_args))
        # reject：以 ToolMessage 回传模型继续规划
        return self.reject(tool_result="用户拒绝了本次操作。...")""")
_bullets(s, Inches(7.85), Inches(1.7), Inches(4.95), Inches(4.4), [
    ("不用现成 ConfirmInterruptRail：", "它只有 approve/reject，edit（改参数后执行）需自定义 Rail"),
    ("条件审批：", "仅当会话已有 html_report 才中断（generate_report 覆盖确认场景）"),
    ("恢复链路：", "InteractiveInput().update(interrupt_id, payload) → 同 session 再次 run_agent"),
    ("审批文案双层：", "框架固定模板 + LLM 生成口语化副标题（5s 超时回退模板）"),
], size=12, space_after=8)
value_bar(s, "approve / edit / reject 三决策用 BaseInterruptRail.resolve_interrupt 严格复刻，拒绝以 ToolMessage 回传而非抛异常",
          y=Inches(6.35))

# ══════════════════════════════════════════════════════════
# 6. 会话与持久化
# ══════════════════════════════════════════════════════════
s = prs.slides.add_slide(blank)
header(s, "应用要点 ③ · 会话与持久化", "框架自带 PersistenceCheckpointer —— 不必自研 checkpointer")
footer(s, page(6))
_card(s, Inches(0.55), Inches(1.6), Inches(6.05), Inches(4.45), "1", "接线方式", [
    "CheckpointerFactory.create(type=\"persistence\", conf={db_type: sqlite, db_timeout: 5, db_enable_wal: True})",
    "Runner.start 之前 set_default_checkpointer（全局默认）",
    "conf 预留 MySQL 通道（db_client 传预构 AsyncEngine）",
    "AgentSession 实例一次性：每轮 run 新建，状态连续性靠 checkpointer 按 session_id 恢复",
    "card 必传，否则 pre_run 报 AttributeError",
], size=12, space_after=7, line_spacing=1.15)
_card(s, Inches(6.85), Inches(1.6), Inches(5.93), Inches(4.45), "2", "分层持久化（核心原则）", [
    "checkpoint blob（pickle）只有框架自己读写 —— 管续聊状态",
    "应用层重载只读业务表（conversation_turn / report_versions）—— 管历史展示，v2/v3 通吃",
    "实测：把 checkpoints.db 整个移走，重载仍返回完整消息",
    "风险：pickle 绑定框架内部结构，升级 openjiuwen 须重验（容错为「上下文重来」，展示不受影响）",
    "审批挂起中重启：作废残留中断 + 修补悬空 tool_calls（否则吞消息 / OpenAI 协议 400）",
], size=12, space_after=7, line_spacing=1.15)
value_bar(s, "checkpoint 管续聊、业务表管展示 —— 两条持久化通道各司其职，互不读取", y=Inches(6.35))

# ══════════════════════════════════════════════════════════
# 7. 流式输出与 SSE 桥接
# ══════════════════════════════════════════════════════════
s = prs.slides.add_slide(blank)
header(s, "应用要点 ④ · 流式输出 → SSE 桥接", "Runner.run_agent_streaming + stream_modes=[OUTPUT, CUSTOM]")
footer(s, page(7))
_table(s, Inches(0.55), Inches(1.6), Inches(7.4), Inches(3.6), [
    ["框架帧类型", "含义", "桥接为"],
    ["llm_output", "正文 token 增量", "正文增量帧"],
    ["llm_reasoning", "思考增量", "思考帧（可折叠）"],
    ["answer", "最终聚合完成帧", "done 帧（error 必须 raise）"],
    ["__interaction__", "中断 / 审批请求", "approval_required 审批卡片"],
    ["llm_usage / tracer", "用量 / 链路追踪", "落库 / 忽略"],
    ["CustomSchema", "工具进度 / 状态更新", "node_update / state_update"],
], col_widths=[Inches(1.9), Inches(2.5), Inches(3.0)])
_bullets(s, Inches(8.25), Inches(1.7), Inches(4.55), Inches(4.4), [
    ("CUSTOM 模式必须显式传：", "否则工具内 write_custom_stream 的进度帧出不来"),
    ("断线重连：", "_BufferedRun 帧缓存 + 游标重放（?after=N）；服务端权威 pending_approval 供刷新对账"),
    ("取消：", "泵任务 + 0.5s 轮询，长 LLM 调用无帧期间也能生效；不要依赖流间隙检查"),
    ("answer 帧陷阱：", "result_type==\"error\" 是框架吞掉的异常，必须主动 raise"),
], size=12, space_after=8)
value_bar(s, "帧映射表按实测类型实现；取消 / 重连 / 审批对账三个机制保住长任务交互体验", y=Inches(6.35))

# ══════════════════════════════════════════════════════════
# 8. 工具改造模式
# ══════════════════════════════════════════════════════════
s = prs.slides.add_slide(blank)
header(s, "应用要点 ⑤ · 业务工具改造：壳 / impl 分层", "本次迁移最值得借鉴的模式 —— 换框架只换壳")
footer(s, page(8))
_code_box(s, Inches(0.55), Inches(1.6), Inches(7.0), Inches(4.5), """@tool
async def search_activities(keyword: str = "", ...) -> str:
    \"\"\"按关键词/渠道/日期等条件查询银行营销活动...\"\"\"
    # 壳只做三件事：组装 ToolContext → 调 impl → 写回 state
    return await _run(impl_activity_query.search_activities,
                      dict(...))

# impl/ 下的业务函数：零框架依赖，可单测、可复用
# 长工具放宽超时：
_t.card.properties["resilience"]["timeout_s"] = 1200""")
_bullets(s, Inches(7.85), Inches(1.7), Inches(4.95), Inches(4.4), [
    ("敏感参数不入模型：", "user_id / org_no / channel / file_path 不在工具签名里，走 ContextVar 双通道注入"),
    ("docstring 即工具描述：", "在首个空行截断 —— 一律无空行单段"),
    ("默认工具超时 300s：", "长工具必须调 resilience.timeout_s（硬上限 3600s），否则管线变孤儿进程"),
    ("impl 层零框架依赖：", "从 langchain 迁来的成本大头因此被削掉"),
], size=12, space_after=8)
value_bar(s, "@tool 壳只管框架交互，业务实现零框架依赖 —— 框架可替换性来自这一层", y=Inches(6.35))

# ══════════════════════════════════════════════════════════
# 9. 迁移方法论
# ══════════════════════════════════════════════════════════
s = prs.slides.add_slide(blank)
header(s, "迁移方法论 · 先 spike，分阶段，每阶段实机验收", "把 0.x 框架的不确定性关在笼子里")
footer(s, page(9))
stages = [
    ("P0", "spike 验证", "只验三件事：端点 tool calling / HITL 全流程 / 流式桥接。任何一条不过即中止", "1~2 天，全部 PASS"),
    ("P1", "替换对话运行时", "agent_v3 并行新建不动 v2；双开关随时回退", "12 项联调全过，5 bug 修复"),
    ("P2", "持久化", "PersistenceCheckpointer 接线 + 业务表重载重写", "重启续聊实测通过"),
    ("P3", "摘除旧依赖", "删 langgraph/langchain/deepagents 全系 8 条依赖，LLM 层也迁", "grep 零命中，终审再抓 3 问题"),
    ("P4", "离线打包", "全量闭包 wheel 集 + 断网安装验证", "断网安装+冒烟全过，包 260M"),
]
y = Inches(1.62)
for i, (tag, name, desc, result) in enumerate(stages):
    chip = _rect(s, Inches(0.55), y, Inches(0.85), Inches(0.78), RED, MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
    _shape_text(chip, tag, size=15, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    tb, tf = _tb(s, Inches(1.65), y + Inches(0.06), Inches(2.1), Inches(0.6))
    set_run(tf.paragraphs[0].add_run(), name, size=14.5, bold=True, color=INK)
    tb, tf = _tb(s, Inches(3.85), y + Inches(0.02), Inches(5.6), Inches(0.85))
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    set_run(tf.paragraphs[0].add_run(), desc, size=11.5, color=GRAY)
    res = _rect(s, Inches(9.65), y + Inches(0.14), Inches(3.13), Inches(0.5), BOX_BG, MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.2)
    _shape_text(res, result, size=10.5, color=INK, align=PP_ALIGN.CENTER)
    y += Inches(0.95)
value_bar(s, "并行新建 + 开关切换，旧链路随时可回退；全部验收过了再删旧代码", y=Inches(6.42), label="关键实践")

# ══════════════════════════════════════════════════════════
# 10. 坑与经验（一）Agent / 工具
# ══════════════════════════════════════════════════════════
s = prs.slides.add_slide(blank)
header(s, "踩坑实录 ① · Agent 配置与工具", "全部实机验证过，源码注释中有对应标记")
footer(s, page(10))
_card(s, Inches(0.55), Inches(1.6), Inches(6.05), Inches(4.45), "1", "Agent 配置", [
    "max_iterations 默认 5 太小 → 显式调大（本项目 50）",
    "AgentSession 实例一次性，复用报 StreamQueue is already closed → 每轮新建 + checkpointer 续状态",
    "import openjiuwen 时 load_dotenv(override=True) 会把 cwd 的 .env 覆盖写进 os.environ → 测试须 delenv 密封",
    "组装时 card 必传，否则 pre_run 报 AttributeError",
    "重复注册同名工具打 ERROR 日志 → 进程级单例只注册一次",
], size=12, space_after=7, line_spacing=1.15)
_card(s, Inches(6.85), Inches(1.6), Inches(5.93), Inches(4.45), "2", "工具", [
    "单次工具调用默认 300s 超时（DEFAULT_TOOL_CALL_TIMEOUT）→ 长工具设 resilience.timeout_s",
    "@tool 的 description 在 docstring 首个空行截断 → 无空行单段",
    "每轮工具调用上限无内置配置 → 自定义 Rail 计数 + _skip_tool + ToolMessage 收尾",
    "模型收到「超限」ToolMessage 后会重试 → skip 文案必须明确「不要再调用」",
    "好消息：from __future__ import annotations 下 @tool 正常（langchain 的坑这里不存在）",
], size=12, space_after=7, line_spacing=1.15)
value_bar(s, "默认值不可信：迭代上限、超时、dotenv 三处默认行为都咬过人", y=Inches(6.35))

# ══════════════════════════════════════════════════════════
# 11. 踩坑实录（二）HITL / 模型 / 持久化 / 流式
# ══════════════════════════════════════════════════════════
s = prs.slides.add_slide(blank)
header(s, "踩坑实录 ② · HITL / 模型层 / 持久化 / 流式", "最难的 bug 出在「审批挂起中重启服务」")
footer(s, page(11))
_card(s, Inches(0.55), Inches(1.6), Inches(6.05), Inches(4.45), "1", "HITL 与流式", [
    "reject 后模型会重试同一工具 → 二次审批：前端状态机必须支持同一轮多次审批",
    "审批挂起中重启会吞新消息：残留中断现场把新消息塞进 INTERACTIVE_INPUT → fresh run 前清残留 + 补合成 ToolMessage（否则 OpenAI 协议 400）",
    "__interaction__ 帧的请求内容在 InteractionOutput.value，不是 .payload",
    "提前 break 消费流触发 ContextVar 警告 → 完整消费 + 泵取消",
], size=11.5, space_after=6, line_spacing=1.12)
_card(s, Inches(6.85), Inches(1.6), Inches(5.93), Inches(4.45), "2", "模型层与持久化", [
    "共享 HTTP 客户端缓存键不含 timeout：主 agent 先建 120s 客户端，管线 600s 超时被静默降为 120s → 逐请求 invoke(timeout=...)",
    "ModelRequestConfig 字段是 model_name（model 仅为 alias）",
    "底层异常被包成 FrameworkError，原始 openai 异常在 __cause__ 链上",
    "agent state 整体 pickle：升级 openjiuwen 版本旧 blob 可能反序列化失败 → 钉版 + 升级重验",
    "CheckpointerConfig 不在 checkpointer 包导出里，要从子模块导入",
], size=11.5, space_after=6, line_spacing=1.12)
value_bar(s, "「审批跨重启恢复」框架其实自带 —— 实测可恢复可执行，未来启用只需产品决策", y=Inches(6.35), label="意外收获")

# ══════════════════════════════════════════════════════════
# 12. 踩坑实录（三）离线交付
# ══════════════════════════════════════════════════════════
s = prs.slides.add_slide(blank)
header(s, "踩坑实录 ③ · 依赖与离线交付", "银行内网场景的最后一公里")
footer(s, page(12))
_card(s, Inches(0.55), Inches(1.6), Inches(6.05), Inches(4.45), "1", "依赖与打包", [
    "依赖 80 项（transformers / pymilvus 等）：既定决策全量闭包不裁剪",
    "cryptography 50.0.0 无 macOS x86_64 wheel（源码编译需 Rust）→ 钉版 46.0.7 写入离线约束",
    "--only-binary 遇无 wheel 包：① 换有 wheel 的兼容版本并钉版 ② 万不得已才 sdist（目标机无编译链）",
    "断网安装铁证：容器 --network none 安装 + import 冒烟 = wheel 闭包自洽",
    "产物：260M tar.gz / 161 个 wheel / MANIFEST.sha256 随包",
], size=11.5, space_after=6, line_spacing=1.12)
_card(s, Inches(6.85), Inches(1.6), Inches(5.93), Inches(4.45), "2", "验证脚本防假阳性", [
    "verify.sh 缺 docker run -i：heredoc 没送进容器，脚本空跑仍输出 VERIFY-PASS（假阳性）",
    "-e EXPECTED_ASSET 未 export：前端新鲜度断言被静默跳过（又一次假阳性）",
    "CA 兜底：openjiuwen 严格 SSL 校验走系统 CA，CA 不全的目标机 LLM 调用全部 CERTIFICATE_VERIFY_FAILED → SSL_CERT_FILE=certifi 兜底",
    "删除的配置键必须从旧 .env 清除：残留 AGENT_V3_ENABLED 会让 pydantic-settings 启动失败",
], size=11.5, space_after=6, line_spacing=1.12)
value_bar(s, "断言不可静默跳过 —— 验证脚本本身也要被验证", y=Inches(6.35))

# ══════════════════════════════════════════════════════════
# 13. 可复用模式 + 总结
# ══════════════════════════════════════════════════════════
s = prs.slides.add_slide(blank)
header(s, "带得走的经验", "可复用适配模式清单")
footer(s, page(13))
_card(s, Inches(0.55), Inches(1.6), Inches(6.05), Inches(4.45), "1", "八个可复用模式", [
    "壳/impl 分层：@tool 壳只管框架交互，业务零框架依赖",
    "ContextVar 双通道注入敏感参数，不进 LLM 上下文",
    "自定义 Rail 三件套：审批 / 限额护栏 / 进度帧",
    "鸭子类型包装 Model（日志 + token 守卫）",
    "分层持久化：checkpoint 管续聊，业务表管展示",
    "to_thread + 线程内独立事件循环跑同步管线",
    "_BufferedRun + 帧游标：断线重连 + 状态对账",
    "角色化模型工厂 + 密文 API key",
], size=11.5, space_after=5, line_spacing=1.12)
_card(s, Inches(6.85), Inches(1.6), Inches(5.93), Inches(4.45), "2", "方法论回顾", [
    "先 spike：1~2 天拿一手结论，不过即中止，成本可控",
    "分阶段替换：并行新建 + 开关切换，随时可回退",
    "每阶段实机验收 + 回归测试网兜底",
    "终审再抓一轮（P3 终审抓到 3 个真实问题）",
    "0.x 框架：钉版 + 升级重验，pickle blob 有容错兜底",
    "核心机制迁移不难，难在四处胶水层：审批 edit / 会话重载 / context 注入 / 多模型路由",
], size=11.5, space_after=6, line_spacing=1.12)
value_bar(s, "用工程方法管住 0.x 框架的不确定性：先验证、可回退、有验收", y=Inches(6.35), label="一句话总结")

out = "openJiuwen框架应用实践_培训.pptx"
prs.save(out)
print("已生成:", out)
