"""外呼活动数据加载与查询服务."""
from __future__ import annotations

import logging
from pathlib import Path
from typing import Any

import openpyxl

logger = logging.getLogger("mock.waihu")

_EXCEL_PATH: Path = Path(__file__).parent / "外呼活动 mock 数据.xlsx"
_data_cache: list[dict[str, Any]] | None = None

_DEMO_NAME_PREFIX = "DEMO-"


def _apply_demo_name_prefix(record: dict[str, Any]) -> None:
    """为活动基本信息中的活动名称添加 DEMO- 前缀（幂等）."""
    basic = record.get("basic_info")
    if not isinstance(basic, dict):
        return
    act_name = basic.get("act_name")
    if isinstance(act_name, str) and act_name and not act_name.startswith(_DEMO_NAME_PREFIX):
        basic["act_name"] = _DEMO_NAME_PREFIX + act_name

# Excel 列头与响应字段的映射
_COLUMN_MAP: list[tuple[str, str]] = [
    ("act_no", "act_no"),
    ("活动基本信息", "basic_info"),
    ("外呼执行情况", "execution"),
    ("外呼所有客户产品交易情况", "all_customers_prodtxn"),
    ("联系失败客户产品交易情况", "failed_customers_prodtxn"),
    ("联系成功客户产品交易情况", "success_customers_prodtxn"),
    ("有效通话客户产品交易情况（prodtxn）", "effective_call_customers_prodtxn"),
    ("有效通话客户产品交易TOP5（prodtop5）", "effective_call_prodtop5"),
    ("有效通话客户存款购买资金来源（prodtxn_ly）", "effective_call_deposit_source"),
    ("有效通话客户理财购买资金来源", "effective_call_wealth_source"),
    ("外呼所有客户资金转化情况", "all_customers_fund_conversion"),
    ("联系成功客户资金转化情况", "success_customers_fund_conversion"),
    ("联系失败客户资金转化情况", "failed_customers_fund_conversion"),
    ("有效通话客户资金转化情况", "effective_call_fund_conversion"),
]


def _parse_kv_block(text: str) -> dict[str, Any]:
    """解析 Excel 中存放的类 JSON 键值对文本块.

    样例数据格式为每行一个 ``"key": value``，行与行之间没有逗号分隔，
    因此不能直接使用标准 json.loads，需要逐行解析。
    """
    result: dict[str, Any] = {}
    if not isinstance(text, str):
        return result

    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line or line in ("{", "}"):
            continue

        # 兼容全角冒号
        line = line.replace("：", ":")
        if ":" not in line:
            continue

        key_part, _, value_part = line.partition(":")
        key = key_part.strip().strip('"').strip("'")
        value = value_part.strip()

        # 去掉行尾可能存在的逗号
        if value.endswith(","):
            value = value[:-1].strip()

        result[key] = _coerce_value(value)

    return result


def _coerce_value(value: str) -> Any:
    """将字符串值转换为合适的 Python 类型."""
    if not isinstance(value, str):
        return value

    # 字符串类型：被引号包裹
    if len(value) >= 2 and value[0] == value[-1] == '"':
        return value[1:-1]
    if len(value) >= 2 and value[0] == value[-1] == "'":
        return value[1:-1]

    # 布尔值
    lowered = value.lower()
    if lowered == "true":
        return True
    if lowered == "false":
        return False
    if lowered in ("null", "none", ""):
        return None

    # 数值类型
    try:
        if "." in value:
            return float(value)
        return int(value)
    except ValueError:
        return value


def load_data() -> list[dict[str, Any]]:
    """从 Excel 加载外呼活动数据.

    Returns:
        每条记录是一个 dict，key 为 _COLUMN_MAP 中定义的英文字段名，
        value 为基本类型或解析后的 dict。
    """
    global _data_cache
    if _data_cache is not None:
        return _data_cache

    if not _EXCEL_PATH.exists():
        logger.error("Excel file not found: %s", _EXCEL_PATH)
        _data_cache = []
        return _data_cache

    wb = openpyxl.load_workbook(str(_EXCEL_PATH), data_only=True)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    header_to_field = dict(_COLUMN_MAP)

    rows: list[dict[str, Any]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if all(v is None for v in row):
            continue

        record: dict[str, Any] = {}
        for header, cell_value in zip(headers, row):
            field_name = header_to_field.get(header)
            if field_name is None:
                continue

            if field_name == "act_no":
                record[field_name] = str(cell_value) if cell_value is not None else ""
            else:
                record[field_name] = _parse_kv_block(cell_value)

        _apply_demo_name_prefix(record)
        rows.append(record)

    _data_cache = rows
    logger.info("Loaded %d waihu activity records from %s", len(rows), _EXCEL_PATH.name)
    return rows


def reload_data() -> list[dict[str, Any]]:
    """强制重新加载 Excel 数据."""
    global _data_cache
    _data_cache = None
    return load_data()


def get_activity_by_id(act_no: str) -> dict[str, Any] | None:
    """根据活动编号查询外呼活动详情."""
    data = load_data()
    for record in data:
        if record.get("act_no") == act_no:
            return record
    return None


def get_activities_by_ids(act_nos: list[str]) -> dict[str, Any]:
    """根据活动编号列表批量查询，返回 {act_no: record}，未命中的编号不出现。"""
    wanted = {str(a).strip() for a in act_nos if str(a).strip()}
    if not wanted:
        return {}
    return {
        record.get("act_no"): record
        for record in load_data()
        if record.get("act_no") in wanted
    }


def list_activities() -> list[dict[str, Any]]:
    """返回外呼活动简要列表."""
    data = load_data()
    result = []
    for record in data:
        basic = record.get("basic_info") or {}
        result.append(
            {
                "act_no": record.get("act_no"),
                "act_name": basic.get("act_name"),
                "yewu_type": basic.get("yewu_type"),
                "活动状态": basic.get("活动状态"),
                "活动人数": basic.get("活动人数"),
            }
        )
    return result


def search_activities(keyword: str) -> list[dict[str, Any]]:
    """根据活动名称模糊搜索外呼活动."""
    keyword = keyword.strip().lower()
    activities = list_activities()
    if not keyword:
        return activities
    return [
        activity
        for activity in activities
        if keyword in str(activity.get("act_name") or "").lower()
    ]
