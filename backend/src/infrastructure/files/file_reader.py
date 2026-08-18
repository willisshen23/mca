from __future__ import annotations

import csv
import json
import math
import statistics
from collections import defaultdict
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

try:
    import xlrd
except ImportError:  # pragma: no cover - optional dependency for legacy .xls files
    xlrd = None

OPENPYXL_EXTENSIONS = {".xlsx", ".xlsm", ".xltx", ".xltm"}
XLRD_EXTENSIONS = {".xls"}
ALLOWED_EXTENSIONS = {".csv", *OPENPYXL_EXTENSIONS, *XLRD_EXTENSIONS}


def read_file_summary(file_path: str) -> str:
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    ext = path.suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {ext}. Allowed: {ALLOWED_EXTENSIONS}")

    if ext == ".csv":
        headers, records = _read_csv(path)
    elif ext in XLRD_EXTENSIONS:
        headers, records = _read_xls(path)
    else:
        headers, records = _read_xlsx(path)

    summary = _build_summary(path, ext, headers, records)
    return json.dumps(summary, ensure_ascii=False, indent=2, default=str)


def read_file_records(file_path: str, max_rows: int = 500) -> dict[str, Any]:
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    ext = path.suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {ext}. Allowed: {ALLOWED_EXTENSIONS}")

    if ext == ".csv":
        headers, records = _read_csv(path)
    elif ext in XLRD_EXTENSIONS:
        headers, records = _read_xls(path)
    else:
        headers, records = _read_xlsx(path)

    limited_records = records[: max(max_rows, 1)]
    return {
        "file_name": path.name,
        "file_type": ext,
        "headers": headers,
        "records": limited_records,
        "total_rows": len(records),
    }


def _read_csv(path: Path) -> tuple[list[str], list[dict[str, Any]]]:
    with path.open("r", encoding="utf-8-sig", newline="") as file:
        reader = csv.reader(file)
        rows = list(reader)

    if not rows:
        return [], []

    headers = _normalize_headers(rows[0])
    records = [_row_to_record(headers, row) for row in rows[1:]]
    return headers, records


def _read_xlsx(path: Path) -> tuple[list[str], list[dict[str, Any]]]:
    workbook = load_workbook(filename=path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows_iter = sheet.iter_rows(values_only=True)
        first_row = next(rows_iter, None)
        if not first_row:
            return [], []

        headers = _normalize_headers(list(first_row))
        records = [_row_to_record(headers, list(row)) for row in rows_iter]
        return headers, records
    finally:
        workbook.close()


def _read_xls(path: Path) -> tuple[list[str], list[dict[str, Any]]]:
    if xlrd is None:
        raise ValueError("Reading .xls files requires dependency xlrd. Please install xlrd>=2.0.")

    workbook = xlrd.open_workbook(str(path))
    sheet = workbook.sheet_by_index(0)
    if sheet.nrows == 0:
        return [], []

    headers = _normalize_headers(sheet.row_values(0))
    records = [_row_to_record(headers, sheet.row_values(row_index)) for row_index in range(1, sheet.nrows)]
    return headers, records


def _normalize_headers(raw_headers: list[Any]) -> list[str]:
    headers: list[str] = []
    seen: dict[str, int] = defaultdict(int)
    for index, raw_header in enumerate(raw_headers, start=1):
        name = str(raw_header).strip() if raw_header not in (None, "") else f"column_{index}"
        seen[name] += 1
        if seen[name] > 1:
            name = f"{name}_{seen[name]}"
        headers.append(name)
    return headers


def _row_to_record(headers: list[str], row: list[Any]) -> dict[str, Any]:
    padded = row[: len(headers)] + [None] * max(0, len(headers) - len(row))
    return {header: _normalize_value(value) for header, value in zip(headers, padded)}


def _normalize_value(value: Any) -> Any:
    if isinstance(value, float) and math.isnan(value):
        return None
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.isoformat()
    return value


def _build_summary(path: Path, ext: str, headers: list[str], records: list[dict[str, Any]]) -> dict[str, Any]:
    column_values = {
        header: [record.get(header) for record in records]
        for header in headers
    }
    column_types = {
        header: _infer_column_type(values)
        for header, values in column_values.items()
    }
    numeric_columns = [
        header
        for header, values in column_values.items()
        if _is_numeric_column(values)
    ]

    summary = {
        "file_name": path.name,
        "file_type": ext,
        "rows": len(records),
        "columns": len(headers),
        "column_names": headers,
        "column_dtypes": column_types,
        "missing_values": {
            header: sum(1 for value in values if _is_missing(value))
            for header, values in column_values.items()
        },
        "numeric_columns": numeric_columns,
        "sample_data": records[:5],
        "basic_stats": {},
    }

    for header in numeric_columns:
        numbers = [_to_float(value) for value in column_values[header] if not _is_missing(value)]
        numbers = [number for number in numbers if number is not None]
        summary["basic_stats"][header] = _numeric_stats(numbers)

    return summary


def _infer_column_type(values: list[Any]) -> str:
    present_values = [value for value in values if not _is_missing(value)]
    if not present_values:
        return "empty"
    if all(_is_number(value) for value in present_values):
        return "number"
    if all(isinstance(value, bool) for value in present_values):
        return "boolean"
    if all(isinstance(value, (datetime, date, time)) for value in present_values):
        return "datetime"
    if all(isinstance(value, str) for value in present_values):
        return "string"
    return "mixed"


def _is_numeric_column(values: list[Any]) -> bool:
    present_values = [value for value in values if not _is_missing(value)]
    return bool(present_values) and all(_is_number(value) for value in present_values)


def _is_number(value: Any) -> bool:
    return _parse_number(value) is not None


def _to_float(value: Any) -> float | None:
    return _parse_number(value)


def _parse_number(value: Any) -> float | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        number = float(value)
    elif isinstance(value, str):
        text = value.strip()
        if not text:
            return None
        try:
            number = float(text)
        except ValueError:
            return None
    else:
        return None

    if math.isnan(number) or math.isinf(number):
        return None
    return number


def _is_missing(value: Any) -> bool:
    return value is None or value == ""


def _numeric_stats(numbers: list[float]) -> dict[str, float | None]:
    if not numbers:
        return {"min": None, "max": None, "mean": None, "median": None, "std": None}
    return {
        "min": min(numbers),
        "max": max(numbers),
        "mean": statistics.fmean(numbers),
        "median": statistics.median(numbers),
        "std": statistics.stdev(numbers) if len(numbers) > 1 else None,
    }
