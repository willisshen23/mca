from __future__ import annotations

from pathlib import Path


def convert_html_to_docx(html_content: str, output_path: str) -> str:
    from bs4 import BeautifulSoup
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)

    soup = BeautifulSoup(html_content, "html.parser")
    doc = Document()

    title_tag = soup.find("h1")
    if title_tag:
        heading = doc.add_heading(title_tag.get_text(strip=True), level=1)
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

    body = soup.find("body") or soup
    for element in body.find_all(
        ["h1", "h2", "h3", "h4", "p", "table", "ul", "ol", "li"],
        recursive=True,
    ):
        tag = element.name
        text = element.get_text(strip=True)

        if not text:
            continue

        if tag in ("h1", "h2", "h3", "h4"):
            doc.add_heading(text, level=int(tag[1]))
        elif tag == "p":
            doc.add_paragraph(text)
        elif tag == "li":
            doc.add_paragraph(text, style="List Bullet")
        elif tag == "table":
            rows = element.find_all("tr")
            if rows:
                headers = [
                    th.get_text(strip=True)
                    for th in rows[0].find_all(["th", "td"])
                ]
                table = doc.add_table(rows=len(rows), cols=len(headers))
                table.style = "Light Grid Accent 1"
                for row_index, row in enumerate(rows):
                    cells = row.find_all(["td", "th"])
                    for column_index, cell in enumerate(cells):
                        if column_index < len(headers):
                            table.cell(row_index, column_index).text = cell.get_text(
                                strip=True
                            )

    doc.save(str(path))
    return str(path.resolve())


def convert_html(html_content: str, target_format: str, output_path: str) -> str:
    if target_format == "docx":
        return convert_html_to_docx(html_content, output_path)
    raise ValueError(f"Unsupported format: {target_format}. Supported: docx")
