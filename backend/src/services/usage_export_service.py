from __future__ import annotations

import io
import json
from datetime import datetime, time, timedelta, timezone

from sqlalchemy import desc
from sqlalchemy.orm import Session

from src.infrastructure.db.models import (
    AgentRunHistory,
    ConversationTurn,
    ReportVersion,
    UsageEvent,
)

EXPORT_MAX_ROWS = 10000

_FEEDBACK_TEXT = {"like": "赞", "dislike": "踩"}


def _period(days: int) -> tuple[datetime, datetime]:
    now = datetime.now(timezone.utc).replace(tzinfo=None)
    end = datetime.combine(now.date() + timedelta(days=1), time.min)
    return end - timedelta(days=days), end


def _format_datetime(value: datetime | None) -> str:
    if value is None:
        return ""
    if value.tzinfo is None:
        value = value.replace(tzinfo=timezone.utc)
    return value.astimezone().strftime("%Y-%m-%d %H:%M:%S")


def _json_list(value: str | None) -> list[str]:
    if not value:
        return []
    try:
        parsed = json.loads(value)
    except (TypeError, ValueError):
        return []
    if not isinstance(parsed, list):
        return []
    return [str(item or "") for item in parsed]


def _template_name(history: AgentRunHistory) -> str:
    snapshot = history.activity_template_snapshot or {}
    values = [
        str(snapshot.get("name") or "").strip(),
        str(history.matched_data_template_name or "").strip(),
    ]
    return " / ".join(dict.fromkeys(value for value in values if value))


def _latest_user_context(
    db: Session,
    user_ids: set[str],
    report_rows: list[tuple[ReportVersion, AgentRunHistory]],
    turn_rows: list[ConversationTurn],
) -> tuple[dict[str, str], dict[str, str]]:
    names: dict[str, str] = {}
    organizations: dict[str, str] = {}

    for _, history in report_rows:
        user_id = str(history.user_id or "")
        if user_id and history.user_name:
            names.setdefault(user_id, history.user_name)
        if user_id and history.org_no:
            organizations.setdefault(user_id, history.org_no)
    for turn in turn_rows:
        user_id = str(turn.user_id or "")
        if user_id and turn.user_name:
            names.setdefault(user_id, turn.user_name)
        if user_id and turn.org_no:
            organizations.setdefault(user_id, turn.org_no)

    if not user_ids:
        return names, organizations

    usage_rows = (
        db.query(UsageEvent.user_id, UsageEvent.org_no)
        .filter(
            UsageEvent.user_id.in_(user_ids),
            UsageEvent.org_no.is_not(None),
            UsageEvent.org_no != "",
        )
        .order_by(desc(UsageEvent.created_at), desc(UsageEvent.id))
        .limit(EXPORT_MAX_ROWS * 5)
        .all()
    )
    for user_id, org_no in usage_rows:
        organizations.setdefault(str(user_id or ""), str(org_no or ""))

    missing_names = user_ids - set(names)
    if missing_names:
        history_names = (
            db.query(AgentRunHistory.user_id, AgentRunHistory.user_name)
            .filter(
                AgentRunHistory.user_id.in_(missing_names),
                AgentRunHistory.user_name.is_not(None),
                AgentRunHistory.user_name != "",
            )
            .order_by(desc(AgentRunHistory.started_at), desc(AgentRunHistory.id))
            .limit(EXPORT_MAX_ROWS)
            .all()
        )
        for user_id, user_name in history_names:
            names.setdefault(str(user_id or ""), str(user_name or ""))

    missing_names = user_ids - set(names)
    if missing_names:
        turn_names = (
            db.query(ConversationTurn.user_id, ConversationTurn.user_name)
            .filter(
                ConversationTurn.user_id.in_(missing_names),
                ConversationTurn.user_name.is_not(None),
                ConversationTurn.user_name != "",
            )
            .order_by(desc(ConversationTurn.started_at), desc(ConversationTurn.id))
            .limit(EXPORT_MAX_ROWS)
            .all()
        )
        for user_id, user_name in turn_names:
            names.setdefault(str(user_id or ""), str(user_name or ""))

    return names, organizations


def _append_sheet(workbook, title: str, headers: list[str], rows: list[list[object]]) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    sheet = workbook.create_sheet(title=title)
    sheet.append(headers)
    for row in rows:
        sheet.append(row)

    header_fill = PatternFill("solid", fgColor="165DFF")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column_cells in sheet.columns:
        max_length = max(
            (len(str(cell.value or "")) for cell in column_cells),
            default=8,
        )
        sheet.column_dimensions[column_cells[0].column_letter].width = min(
            max(max_length + 3, 12),
            36,
        )


def export_usage_details_xlsx(db: Session, days: int) -> bytes:
    from openpyxl import Workbook

    start, end = _period(days)
    login_rows = (
        db.query(UsageEvent)
        .filter(
            UsageEvent.event_type == "home_view",
            UsageEvent.created_at >= start,
            UsageEvent.created_at < end,
        )
        .order_by(desc(UsageEvent.created_at), desc(UsageEvent.id))
        .limit(EXPORT_MAX_ROWS)
        .all()
    )
    report_rows = (
        db.query(ReportVersion, AgentRunHistory)
        .join(AgentRunHistory, AgentRunHistory.request_id == ReportVersion.request_id)
        .filter(
            ReportVersion.source == "generate",
            ReportVersion.created_at >= start,
            ReportVersion.created_at < end,
        )
        .order_by(desc(ReportVersion.created_at), desc(ReportVersion.id))
        .limit(EXPORT_MAX_ROWS)
        .all()
    )
    turn_rows = (
        db.query(ConversationTurn)
        .filter(
            ConversationTurn.started_at >= start,
            ConversationTurn.started_at < end,
        )
        .order_by(desc(ConversationTurn.started_at), desc(ConversationTurn.id))
        .limit(EXPORT_MAX_ROWS)
        .all()
    )

    user_ids = {
        str(user_id)
        for user_id in [
            *(row.user_id for row in login_rows),
            *(history.user_id for _, history in report_rows),
            *(row.user_id for row in turn_rows),
        ]
        if user_id
    }
    names, organizations = _latest_user_context(db, user_ids, report_rows, turn_rows)

    login_data = [
        [
            row.user_id or "",
            "活动分析助手",
            names.get(row.user_id or "", ""),
            row.org_no or organizations.get(row.user_id or "", ""),
            _format_datetime(row.created_at),
        ]
        for row in login_rows
    ]
    report_data = [
        [
            history.user_id or "",
            history.user_name or names.get(history.user_id or "", ""),
            history.org_no or organizations.get(history.user_id or "", ""),
            _template_name(history),
            _format_datetime(version.created_at),
            _FEEDBACK_TEXT.get(version.feedback or history.report_feedback or "", ""),
            round(history.duration_ms / 1000, 1)
            if history.duration_ms is not None
            else "",
        ]
        for version, history in report_rows
    ]
    conversation_data = [
        [
            row.user_id or "",
            row.user_name or names.get(row.user_id or "", ""),
            row.org_no or organizations.get(row.user_id or "", ""),
            "是" if row.status == "success" else ("否" if row.status else ""),
            _format_datetime(row.started_at),
        ]
        for row in turn_rows
    ]

    activity_data: list[list[object]] = []
    for version, history in report_rows:
        activity_ids = _json_list(history.activity_ids_json)
        activity_names = _json_list(history.activity_names_json)
        count = max(len(activity_ids), len(activity_names), 1)
        for index in range(count):
            activity_data.append(
                [
                    activity_ids[index] if index < len(activity_ids) else "",
                    activity_names[index] if index < len(activity_names) else "",
                    version.report_asset_name or "",
                    history.user_name or history.user_id or "",
                    history.org_no or organizations.get(history.user_id or "", ""),
                ]
            )

    workbook = Workbook()
    workbook.remove(workbook.active)
    _append_sheet(
        workbook,
        "登陆情况",
        ["员工编号", "登陆系统", "员工姓名", "员工所属部门", "登陆时间"],
        login_data,
    )
    _append_sheet(
        workbook,
        "报告生成情况",
        [
            "员工编号",
            "员工姓名",
            "员工所属部门",
            "报告模板",
            "生成时间",
            "赞/踩",
            "报告耗时（秒）",
        ],
        report_data,
    )
    _append_sheet(
        workbook,
        "会话情况",
        ["员工编号", "员工姓名", "员工所属部门", "会话是否成功", "会话创建时间"],
        conversation_data,
    )
    _append_sheet(
        workbook,
        "活动分析情况",
        ["活动编号", "活动名称", "活动报告", "报告创建人", "报告创建人所属部门"],
        activity_data,
    )

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
