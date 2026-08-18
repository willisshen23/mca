"""报告版本链（report_versions）服务层测试。

覆盖设计文档 §7 pytest 清单：
save_report_file 版本递增/change_note/latest_version_no 同步、反馈按版本、
会话聚合与筛选分页、详情版本摘要、版本下载与遗留兜底、迁移幂等、版本粒度导出。
"""

from __future__ import annotations

import io
from datetime import datetime, timedelta, timezone

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine, text
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

import src.services.agent_run_history_service as svc
from src.infrastructure.db.models import (
    AgentRunHistory,
    Base,
    ConversationMeta,
    ReportVersion,
)


@pytest.fixture
def Session(monkeypatch):
    engine = create_engine(
        "sqlite:///:memory:",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(bind=engine)
    session_factory = sessionmaker(
        bind=engine, autoflush=False, autocommit=False, expire_on_commit=False
    )
    monkeypatch.setattr(svc, "SessionLocal", session_factory)
    monkeypatch.setattr(svc, "engine", engine)
    return session_factory


def _add_history(Session, request_id="r1", thread_id="t1", **overrides):
    defaults = {
        "request_id": request_id,
        "thread_id": thread_id,
        "user_id": "u1",
        "user_name": "张三",
        "user_text": "帮我分析一下活动",
        "status": "success",
        "prompt_text": "prompt",
        "started_at": datetime.now(timezone.utc),
    }
    defaults.update(overrides)
    with Session() as s:
        row = AgentRunHistory(**defaults)
        s.add(row)
        s.commit()
        s.refresh(row)
        return row.id


def _get_version(Session, request_id, version_no):
    with Session() as s:
        return (
            s.query(ReportVersion)
            .filter_by(request_id=request_id, version_no=version_no)
            .first()
        )


# ── save_report_file ──


def test_save_report_file_creates_sequential_versions(Session):
    _add_history(Session)

    asset_name, version_no = svc.save_report_file("r1", "<html><body>v1</body></html>")
    assert asset_name == "activity-analysis-report-r1.html"
    assert version_no == 1

    _, version_no = svc.save_report_file(
        "r1", "<html><body>v2</body></html>", source="modify", change_note="把标题改短"
    )
    assert version_no == 2

    _, version_no = svc.save_report_file(
        "r1", "<html><body>v3</body></html>", source="modify", change_note="x" * 600
    )
    assert version_no == 3

    with Session() as s:
        row = s.query(AgentRunHistory).filter_by(request_id="r1").one()
        assert row.latest_version_no == 3
        # 旧列停止写入
        assert row.report_asset_blob is None
        assert row.report_asset_name is None
        versions = (
            s.query(ReportVersion)
            .filter_by(request_id="r1")
            .order_by(ReportVersion.version_no)
            .all()
        )
    assert [v.version_no for v in versions] == [1, 2, 3]
    assert versions[0].source == "generate"
    assert versions[0].change_note is None
    assert versions[1].source == "modify"
    assert versions[1].change_note == "把标题改短"
    assert len(versions[2].change_note) == 500
    assert b"v2" in bytes(versions[1].report_asset_blob)


def test_save_report_file_failure_returns_empty(Session):
    assert svc.save_report_file("missing", "<html>v1</html>") == ("", None)
    _add_history(Session)
    assert svc.save_report_file("r1", "") == ("", None)


# ── 反馈 ──


def _seed_two_versions(Session, request_id="r1", thread_id="t1"):
    _add_history(Session, request_id=request_id, thread_id=thread_id)
    svc.save_report_file(request_id, "<html>v1</html>")
    svc.save_report_file(request_id, "<html>v2</html>", source="modify", change_note="改")


def test_feedback_targets_specific_version(Session):
    _seed_two_versions(Session)

    result = svc.update_report_feedback("r1", "like", "u1", version_no=1)
    # 响应必须反映实际被反馈的版本，而非最新版本（否则前端卡片版本号/反馈显示错乱）
    assert result["version_no"] == 1
    assert result["report_feedback"] == "like"
    v1 = _get_version(Session, "r1", 1)
    v2 = _get_version(Session, "r1", 2)
    assert v1.feedback == "like"
    assert v1.feedback_user_id == "u1"
    assert v1.feedback_at is not None
    assert v2.feedback is None


def test_feedback_without_version_targets_latest(Session):
    _seed_two_versions(Session)
    svc.update_report_feedback("r1", "dislike", "u2", comment="数据不对")
    v1 = _get_version(Session, "r1", 1)
    v2 = _get_version(Session, "r1", 2)
    assert v1.feedback is None
    assert v2.feedback == "dislike"
    assert v2.feedback_comment == "数据不对"

    result = svc.update_report_feedback("r1", "dislike", "u2", comment="数据不对")
    assert result["report_feedback"] == "dislike"
    assert result["feedback_comment"] == "数据不对"


def test_feedback_comment_only_kept_for_dislike(Session):
    _seed_two_versions(Session)
    svc.update_report_feedback("r1", "like", "u1", comment="顺手写的")
    v2 = _get_version(Session, "r1", 2)
    assert v2.feedback == "like"
    assert v2.feedback_comment is None


def test_feedback_none_clears_version_feedback(Session):
    _seed_two_versions(Session)
    svc.update_report_feedback("r1", "dislike", "u1", comment="原因")
    svc.update_report_feedback("r1", "none")
    v2 = _get_version(Session, "r1", 2)
    assert v2.feedback is None
    assert v2.feedback_comment is None
    assert v2.feedback_user_id is None
    assert v2.feedback_at is None


def test_feedback_invalid_value_raises(Session):
    _seed_two_versions(Session)
    with pytest.raises(ValueError):
        svc.update_report_feedback("r1", "awesome")
    with pytest.raises(ValueError):
        svc.update_report_feedback("r1", "dislike", comment="x" * 501)


def test_feedback_unknown_request_or_version_returns_none(Session):
    _seed_two_versions(Session)
    assert svc.update_report_feedback("nope", "like") is None
    assert svc.update_report_feedback("r1", "like", version_no=99) is None


def test_feedback_by_thread_targets_latest_report_latest_version(Session):
    now = datetime.now(timezone.utc)
    _seed_two_versions(Session, request_id="r-old", thread_id="t1")
    _add_history(Session, request_id="r-new", thread_id="t1", started_at=now + timedelta(minutes=1))
    svc.save_report_file("r-new", "<html>new v1</html>")
    svc.save_report_file("r-new", "<html>new v2</html>", source="modify", change_note="改")

    result = svc.update_latest_report_feedback_by_thread("t1", "like", "u1")
    assert result["request_id"] == "r-new"
    assert result["version_no"] == 2
    assert _get_version(Session, "r-new", 2).feedback == "like"
    assert _get_version(Session, "r-old", 2).feedback is None
    assert svc.update_latest_report_feedback_by_thread("t-unknown", "like") is None


# ── 会话聚合列表 ──


def _seed_sessions(Session):
    now = datetime.now(timezone.utc)
    # 会话 t1：两份报告，v1 赞 / 另一份踩
    _seed_two_versions(Session, request_id="r1", thread_id="t1")
    _add_history(Session, request_id="r2", thread_id="t1", started_at=now + timedelta(minutes=5))
    svc.save_report_file("r2", "<html>r2v1</html>")
    svc.update_report_feedback("r1", "like", "u1", version_no=1)
    svc.update_report_feedback("r2", "dislike", "u1", comment="不好")
    # 会话 t2：一份报告，无反馈
    _add_history(Session, request_id="r3", thread_id="t2", user_text="单独的提问", started_at=now + timedelta(minutes=10))
    svc.save_report_file("r3", "<html>r3v1</html>")

    with Session() as s:
        s.add(ConversationMeta(thread_id="t1", user_id="u1", title="会话A"))
        s.commit()


def test_list_report_sessions_aggregates(Session):
    _seed_sessions(Session)
    result = svc.list_report_sessions()
    assert result["total"] == 2
    items = {item["thread_id"]: item for item in result["items"]}
    t1 = items["t1"]
    assert t1["title"] == "会话A"
    assert t1["report_count"] == 2
    assert t1["version_count"] == 3
    assert t1["like_count"] == 1
    assert t1["dislike_count"] == 1
    t2 = items["t2"]
    assert t2["title"] == "单独的提问"  # 无 conversation_meta 回退最早 user_text
    assert t2["report_count"] == 1
    assert t2["version_count"] == 1
    assert t2["like_count"] == 0
    assert t2["dislike_count"] == 0
    # last_active_at 倒序
    assert result["items"][0]["thread_id"] == "t2"


def test_list_report_sessions_filters(Session):
    _seed_sessions(Session)
    assert svc.list_report_sessions(keyword="会话A")["total"] == 1
    assert svc.list_report_sessions(keyword="单独的提问")["total"] == 1
    assert svc.list_report_sessions(keyword="r2")["total"] == 1
    assert svc.list_report_sessions(keyword="不存在")["total"] == 0

    dislike_only = svc.list_report_sessions(feedback="dislike")
    assert dislike_only["total"] == 1
    assert dislike_only["items"][0]["thread_id"] == "t1"
    none_only = svc.list_report_sessions(feedback="none")
    assert none_only["total"] == 1
    assert none_only["items"][0]["thread_id"] == "t2"

    today = datetime.now(timezone.utc).strftime("%Y-%m-%d")
    assert svc.list_report_sessions(date_from=today, date_to=today)["total"] == 2
    future = (datetime.now(timezone.utc) + timedelta(days=7)).strftime("%Y-%m-%d")
    assert svc.list_report_sessions(date_from=future)["total"] == 0
    with pytest.raises(ValueError):
        svc.list_report_sessions(date_from="2026/01/01")


def test_list_report_sessions_pagination(Session):
    _seed_sessions(Session)
    page1 = svc.list_report_sessions(limit=1, offset=0)
    page2 = svc.list_report_sessions(limit=1, offset=1)
    assert page1["total"] == 2 and len(page1["items"]) == 1
    assert len(page2["items"]) == 1
    assert page1["items"][0]["thread_id"] != page2["items"][0]["thread_id"]


# ── 会话详情 ──


def test_get_report_session_detail(Session):
    _seed_sessions(Session)
    detail = svc.get_report_session_detail("t1")
    assert detail is not None
    assert detail["title"] == "会话A"
    assert detail["report_count"] == 2
    assert detail["version_count"] == 3
    assert detail["like_count"] == 1
    assert detail["dislike_count"] == 1
    assert [report["request_id"] for report in detail["reports"]] == ["r2", "r1"]

    r1_report = next(r for r in detail["reports"] if r["request_id"] == "r1")
    assert r1_report["latest_version_no"] == 2
    assert [v["version_no"] for v in r1_report["versions"]] == [1, 2]
    v1 = r1_report["versions"][0]
    assert v1["source"] == "generate"
    assert v1["feedback"] == "like"
    v2 = r1_report["versions"][1]
    assert v2["source"] == "modify"
    assert v2["change_note"] == "改"
    # 版本摘要不含 blob
    for version in r1_report["versions"]:
        assert "report_asset_blob" not in version

    assert svc.get_report_session_detail("t-unknown") is None


# ── 下载 ──


def test_get_report_asset_default_latest_and_specific_version(Session):
    history_id = _add_history(Session)
    svc.save_report_file("r1", "<html><body>v1-content</body></html>")
    svc.save_report_file("r1", "<html><body>v2-content</body></html>", source="modify", change_note="改")

    latest = svc.get_report_asset(history_id)
    assert latest is not None
    assert b"v2-content" in latest.content

    v1 = svc.get_report_asset(history_id, version_no=1)
    assert v1 is not None
    assert b"v1-content" in v1.content

    assert svc.get_report_asset(history_id, version_no=99) is None
    assert svc.get_report_asset(99999) is None


def test_get_report_asset_legacy_fallback(Session):
    legacy_html = "<html><body>legacy</body></html>".encode("utf-8")
    history_id = _add_history(
        Session,
        report_asset_name="legacy.html",
        report_asset_media_type="text/html; charset=utf-8",
        report_asset_blob=legacy_html,
    )
    asset = svc.get_report_asset(history_id)
    assert asset is not None
    assert asset.content == legacy_html
    # 遗留数据无版本行，指定版本号 → 404 语义（None）
    assert svc.get_report_asset(history_id, version_no=1) is None


# ── 迁移 ──


def test_ensure_report_versions_schema_migrates_and_idempotent(Session):
    engine = svc.engine
    legacy_html = "<html><body>legacy blob</body></html>".encode("utf-8")
    _add_history(
        Session,
        request_id="legacy-1",
        report_asset_name="legacy.html",
        report_asset_media_type="text/html; charset=utf-8",
        report_asset_blob=legacy_html,
        report_feedback="dislike",
        feedback_user_id="u9",
    )
    _add_history(Session, request_id="empty-1")  # 无报告，不应产生版本

    # 模拟旧库：无 report_versions 表、无 latest_version_no 列
    with engine.begin() as conn:
        conn.execute(text("DROP TABLE report_versions"))
        conn.execute(text("ALTER TABLE agent_run_history DROP COLUMN latest_version_no"))

    # 第一次：补列；表不存在跳过数据搬迁
    svc.ensure_report_versions_schema()

    # 建表后第二次：搬迁 blob 与反馈
    ReportVersion.__table__.create(bind=engine)
    svc.ensure_report_versions_schema()
    with Session() as s:
        versions = s.query(ReportVersion).all()
        assert len(versions) == 1
        version = versions[0]
        assert version.request_id == "legacy-1"
        assert version.version_no == 1
        assert version.source == "generate"
        assert bytes(version.report_asset_blob) == legacy_html
        assert version.feedback == "dislike"
        assert version.feedback_user_id == "u9"
        row = s.query(AgentRunHistory).filter_by(request_id="legacy-1").one()
        assert row.latest_version_no == 1

    # 第三次：幂等，不重复建行
    svc.ensure_report_versions_schema()
    with Session() as s:
        assert s.query(ReportVersion).count() == 1


# ── 导出 ──


def test_export_report_sessions_xlsx(Session):
    _seed_sessions(Session)
    content = svc.export_report_sessions_xlsx()
    workbook = load_workbook(io.BytesIO(content))
    sheet = workbook.active
    header = [cell.value for cell in sheet[1]]
    assert header == [
        "时间",
        "会话标题",
        "发起人",
        "request_id",
        "版本号",
        "来源",
        "修改说明",
        "分析模板",
        "数据模板",
        "状态",
        "耗时(ms)",
        "反馈",
        "点踩原因",
        "反馈人",
        "反馈时间",
    ]
    data_rows = list(sheet.iter_rows(min_row=2, values_only=True))
    assert len(data_rows) == 4  # r1×2 + r2×1 + r3×1，版本粒度
    by_request = {}
    for row in data_rows:
        by_request.setdefault(row[3], []).append(row)
    assert sorted(v[4] for v in by_request["r1"]) == [1, 2]
    r2_row = by_request["r2"][0]
    assert r2_row[11] == "已点踩"
    assert r2_row[12] == "不好"
    modify_row = next(v for v in by_request["r1"] if v[4] == 2)
    assert modify_row[5] == "修改"
    assert modify_row[6] == "改"

    # 反馈筛选（版本粒度）
    liked = load_workbook(io.BytesIO(svc.export_report_sessions_xlsx(feedback="like"))).active
    assert len(list(liked.iter_rows(min_row=2, values_only=True))) == 1


# ── 序列化 ──


def test_serialize_row_includes_version_fields(Session):
    _seed_two_versions(Session)
    svc.update_report_feedback("r1", "dislike", "u1", comment="原因文本", version_no=2)
    with Session() as s:
        row = s.query(AgentRunHistory).filter_by(request_id="r1").one()
        payload = svc._serialize_row(row, include_prompt=False, db=s)
    assert payload["latest_version_no"] == 2
    assert payload["version_no"] == 2
    assert payload["report_feedback"] == "dislike"
    assert payload["feedback_comment"] == "原因文本"
    assert payload["has_report"] is True


def test_serialize_row_legacy_fallback(Session):
    _add_history(
        Session,
        report_asset_name="legacy.html",
        report_asset_blob=b"<html>legacy</html>",
        report_feedback="like",
    )
    with Session() as s:
        row = s.query(AgentRunHistory).filter_by(request_id="r1").one()
        payload = svc._serialize_row(row, include_prompt=False, db=s)
    assert payload["version_no"] is None
    assert payload["report_feedback"] == "like"  # 无版本行回退旧列
    assert payload["feedback_comment"] == ""
