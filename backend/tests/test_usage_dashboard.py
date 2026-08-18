from __future__ import annotations

import unittest
from datetime import datetime

from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.dialects import mysql
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from src.infrastructure.db.models import (
    AgentRunHistory,
    Base,
    ConversationMeta,
    ConversationTurn,
    LlmTokenUsageLog,
    ReportVersion,
    UsageEvent,
)
from src.infrastructure.db.session import get_db
from src.main import app

engine = create_engine(
    "sqlite:///:memory:",
    connect_args={"check_same_thread": False},
    poolclass=StaticPool,
)
TestingSessionLocal = sessionmaker(bind=engine, autoflush=False, expire_on_commit=False)


def override_get_db():
    db = TestingSessionLocal()
    try:
        yield db
    finally:
        db.close()


class UsageDashboardTest(unittest.TestCase):
    def setUp(self):
        Base.metadata.create_all(bind=engine)
        app.dependency_overrides[get_db] = override_get_db
        self.client = TestClient(app)

    def tearDown(self):
        app.dependency_overrides.pop(get_db, None)
        Base.metadata.drop_all(bind=engine)

    def test_event_collection_and_summary_use_real_records(self):
        headers = {"X-User-Id": "user-001"}
        for event_type in ("home_view", "home_view"):
            response = self.client.post(
                "/api/usage-dashboard/events",
                headers=headers,
                json={
                    "event_type": event_type,
                    "org_no": "320100",
                    "session_id": "session-001",
                    "path": "/management",
                    "source": "portal",
                },
            )
            self.assertEqual(response.status_code, 201)

        with TestingSessionLocal() as db:
            now = datetime.now()
            db.add(
                ConversationMeta(
                    thread_id="thread-001",
                    user_id="user-001",
                    title="测试会话",
                    agent_version="v2",
                )
            )
            db.add(
                ConversationTurn(
                    request_id="turn-001",
                    thread_id="thread-001",
                    user_id="user-001",
                    user_name="张三",
                    org_no="320100",
                    status="success",
                    started_at=now,
                    finished_at=now,
                    duration_ms=2500,
                )
            )
            db.add_all(
                [
                    LlmTokenUsageLog(
                        id="llm-run-001",
                        start_time=now,
                        time_spent=1200,
                        token_cost=2400,
                        exec_status="SUCCESS",
                    ),
                    LlmTokenUsageLog(
                        id="llm-run-002",
                        start_time=now,
                        time_spent=800,
                        token_cost=600,
                        exec_status="FAILED",
                    ),
                ]
            )
            db.add(
                AgentRunHistory(
                    request_id="report-001",
                    thread_id="thread-001",
                    user_id="user-001",
                    user_name="张三",
                    org_no="320100",
                    status="success",
                    started_at=now,
                    finished_at=now,
                    duration_ms=2500,
                    activity_template_snapshot={"name": "营销分析模板"},
                    activity_ids_json='["ACT-001"]',
                    activity_names_json='["夏季营销活动"]',
                )
            )
            db.flush()
            db.add(
                ReportVersion(
                    request_id="report-001",
                    version_no=1,
                    report_asset_name="report.html",
                    report_asset_media_type="text/html",
                    report_asset_blob=b"<html></html>",
                    source="generate",
                    feedback="like",
                )
            )
            db.commit()

        response = self.client.get("/api/usage-dashboard/summary?days=30")
        self.assertEqual(response.status_code, 200)
        data = response.json()
        self.assertEqual(data["metrics"]["visitors"]["value"], 1)
        self.assertEqual(data["metrics"]["page_views"]["value"], 2)
        self.assertEqual(data["metrics"]["organizations"]["value"], 1)
        self.assertEqual(data["metrics"]["conversations"]["value"], 1)
        self.assertEqual(data["metrics"]["turns"]["value"], 1)
        self.assertEqual(data["metrics"]["reports"]["value"], 1)
        self.assertEqual(data["metrics"]["success_rate"], 100.0)
        self.assertEqual(data["metrics"]["avg_duration_seconds"], 2.5)
        self.assertEqual(data["llm_usage"]["calls"]["value"], 2)
        self.assertEqual(data["llm_usage"]["total_tokens"]["value"], 3000)
        self.assertEqual(data["llm_usage"]["avg_tokens_per_call"], 1500.0)
        self.assertEqual(data["llm_usage"]["avg_time_spent_ms"], 1000.0)
        self.assertEqual(data["llm_usage"]["success_rate"], 50.0)
        self.assertEqual(data["llm_usage"]["successful_calls"], 1)
        self.assertEqual(data["llm_usage"]["failed_calls"], 1)
        self.assertEqual(len(data["llm_usage"]["trend"]), 30)
        today_usage = next(
            item
            for item in data["llm_usage"]["trend"]
            if item["date"] == now.date().isoformat()
        )
        self.assertEqual(today_usage["tokens"], 3000)
        self.assertEqual(today_usage["calls"], 2)
        self.assertEqual(
            data["organizations"],
            [{"org_no": "320100", "visitors": 1, "page_views": 2}],
        )
        self.assertEqual([item["key"] for item in data["funnel"]], [
            "visitors",
            "conversations",
            "reports",
        ])
        self.assertEqual(len(data["trend"]), 30)

        export_response = self.client.get("/api/usage-dashboard/export?days=30")
        self.assertEqual(export_response.status_code, 200)
        self.assertIn(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            export_response.headers["content-type"],
        )
        from openpyxl import load_workbook
        from io import BytesIO

        workbook = load_workbook(BytesIO(export_response.content), read_only=True)
        self.assertEqual(
            workbook.sheetnames,
            ["登陆情况", "报告生成情况", "会话情况", "活动分析情况"],
        )
        self.assertEqual(workbook["登陆情况"]["A2"].value, "user-001")
        self.assertEqual(workbook["登陆情况"]["C2"].value, "张三")
        self.assertEqual(workbook["登陆情况"]["D2"].value, "320100")
        self.assertEqual(workbook["报告生成情况"]["A2"].value, "user-001")
        self.assertEqual(workbook["报告生成情况"]["D2"].value, "营销分析模板")
        self.assertEqual(workbook["报告生成情况"]["F2"].value, "赞")
        self.assertEqual(workbook["报告生成情况"]["G2"].value, 2.5)
        self.assertEqual(workbook["会话情况"]["D2"].value, "是")
        self.assertEqual(workbook["活动分析情况"]["A2"].value, "ACT-001")
        self.assertEqual(workbook["活动分析情况"]["B2"].value, "夏季营销活动")

    def test_event_validation_and_period_validation(self):
        response = self.client.post(
            "/api/usage-dashboard/events",
            headers={"X-User-Id": "user-001"},
            json={"event_type": "unknown"},
        )
        self.assertEqual(response.status_code, 422)
        self.assertEqual(
            self.client.get("/api/usage-dashboard/summary?days=14").status_code,
            422,
        )

    def test_usage_event_schema_has_mysql_safe_field_types(self):
        table = UsageEvent.__table__
        dialect = mysql.dialect()
        self.assertEqual(table.c.event_type.type.compile(dialect=dialect), "VARCHAR(40)")
        self.assertEqual(table.c.user_id.type.compile(dialect=dialect), "VARCHAR(64)")
        self.assertEqual(table.c.org_no.type.compile(dialect=dialect), "VARCHAR(64)")
        self.assertEqual(table.c.referrer.type.compile(dialect=dialect), "VARCHAR(500)")
        self.assertEqual(table.c.metadata_json.type.compile(dialect=dialect), "TEXT")
        history_table = AgentRunHistory.__table__
        self.assertEqual(
            history_table.c.org_no.type.compile(dialect=dialect),
            "VARCHAR(64)",
        )
        self.assertEqual(
            history_table.c.activity_ids_json.type.compile(dialect=dialect),
            "TEXT",
        )
        self.assertEqual(
            ConversationTurn.__table__.c.org_no.type.compile(dialect=dialect),
            "VARCHAR(64)",
        )


if __name__ == "__main__":
    unittest.main()
